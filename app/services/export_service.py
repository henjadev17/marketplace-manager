from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def export_products(rows, output_file: Path):
    max_photos = max((len(r["photos"]) for r in rows), default=0)
    headers = [
        "Código",
        "Título",
        "Precio",
        "Descripción específica",
        "Descripción final",
        "Categoría",
        "Ubicación",
        "Estado",
        "Cantidad fotos",
    ] + [f"Foto {i}" for i in range(1, max_photos + 1)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(
            [
                row["code"],
                row["title"],
                row["price"],
                row["description"],
                row.get("final_description", ""),
                row["category"],
                row["location"],
                row["status"],
                len(row["photos"]),
                *row["photos"],
            ]
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for idx, col in enumerate(ws.columns, start=1):
        width = max(
            len(str(cell.value)) if cell.value is not None else 0 for cell in col
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(
            max(width + 2, 12), 65
        )

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    return output_file
